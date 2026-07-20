from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
OUT = Path(__file__).resolve().parent / "EchoesOfTheRuins_CourseworkReport.docx"
NAVY = "0B1A2B"
CYAN = "24D8FF"
SLATE = "334155"
MUTED = "64748B"
PALE = "EAF2F8"


def set_font(run, size=11, color="000000", bold=False, italic=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def set_cell_width(cell, dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    width = tc_pr.first_child_found_in("w:tcW")
    if width is None:
        width = OxmlElement("w:tcW")
        tc_pr.append(width)
    width.set(qn("w:w"), str(dxa))
    width.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    indent = OxmlElement("w:tblInd")
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")
    tbl_pr.append(indent)
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def paragraph(doc, text="", style=None, before=0, after=6, align=None):
    p = doc.add_paragraph(style=style)
    if text:
        set_font(p.add_run(text))
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    if align is not None:
        p.alignment = align
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_font(run, 15 if level == 1 else 12, CYAN if level == 1 else SLATE, bold=True)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.1
    set_font(p.add_run(text), 10.5)


def image(doc, path, width=5.9):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(5)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))


def make_header_footer(section):
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("ECHOES OF THE RUINS  |  COURSEWORK REPORT")
    set_font(r, 8.5, MUTED, bold=True)
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("Unity 6.3 LTS | Windows single-player vertical slice")
    set_font(r, 8, MUTED)


def add_metadata(doc):
    table = doc.add_table(rows=2, cols=2)
    set_table_geometry(table, [2100, 7260])
    rows = [("Project", "Echoes of the Ruins"), ("Engine / target", "Unity 6.3 LTS, URP, Windows 64-bit")]
    for row, (label, value) in zip(table.rows, rows):
        shade(row.cells[0], PALE)
        p = row.cells[0].paragraphs[0]
        set_font(p.add_run(label), 10, NAVY, bold=True)
        p = row.cells[1].paragraphs[0]
        set_font(p.add_run(value), 10)


def add_evidence_table(doc):
    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [2400, 4600, 2360])
    headers = ["Evidence", "What it demonstrates", "Result"]
    for cell, text in zip(table.rows[0].cells, headers):
        shade(cell, PALE)
        set_font(cell.paragraphs[0].add_run(text), 9.5, NAVY, bold=True)
    rows = [
        ("EditMode", "Interaction timing, save data, scoring, NavMesh and scene integrity.", "156 / 156 passed"),
        ("PlayMode", "Core-to-exit flow, guardian states, capture and HUD events.", "14 / 14 passed"),
        ("Windows build", "Windows x64 player with no missing-script, NavMesh, shader or C# errors.", "Build completed"),
    ]
    for left, mid, right in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, (left, mid, right)):
            p = cell.paragraphs[0]
            set_font(p.add_run(text), 9.2)
            p.paragraph_format.space_after = Pt(2)


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    make_header_footer(section)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.5)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(22)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("ECHOES OF THE RUINS")
    set_font(r, 25, NAVY, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run("Coursework report: a moonlit stealth-exploration vertical slice")
    set_font(r, 12.5, SLATE)
    add_metadata(doc)

    heading(doc, "1. Game overview and selling point")
    paragraph(doc, "Echoes of the Ruins is a third-person stealth-exploration game for casual players. The player enters a moonlit citadel, attunes three energy cores and escapes through the north gate while avoiding three stone guardians. The core loop is deliberately readable: observe patrols, use shadows and spend limited echo stones to redirect danger. The project is scoped as one 8-10 minute polished vertical slice rather than several unfinished maps.")

    heading(doc, "2. Story, audience and player character")
    paragraph(doc, "The ruins are waking after centuries of silence. A hooded relic hunter must restore the stolen cores before the guardian network seals the citadel. The explorer silhouette, cyan amulet and shoulder camera remain readable against blue-grey stone. Staged teaching, forgiving suspicion accumulation and checkpoint recovery support casual players; optional relics and score ranks create replay value without requiring combat.")

    heading(doc, "3. Visual and audio direction")
    paragraph(doc, "The visual language uses cold moonlight for architecture, cyan for objectives and warm amber for fire and guardian danger. URP supplies restrained bloom, colour adjustment, fog and high-quality antialiasing. Modular CC0 ruins, masonry walls, broken arches, distant silhouettes, braziers and layered stone textures replace the prototype cubes. Motion blur, temporal antialiasing and film grain are disabled to avoid a smeared moving-floor image. Sound communicates footsteps, core progress, investigation, capture and exit unlock; warnings never rely on colour alone.")
    image(doc, ROOT / "Evidence" / "windows-main-menu.png", 5.75)
    p = paragraph(doc, "Figure 1. Final main menu: explicit objective, New Game, Continue, Controls, Settings and Quit.", after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    for run in p.runs:
        set_font(run, 8.5, MUTED, italic=True)

    heading(doc, "4. Level, controls and first-minute experience")
    paragraph(doc, "The approximately 129-metre route has six zones: Safe Entry, Courtyard, Shadow Gallery, Echo Passage, Altar and North Gate. The first-minute HUD states the whole rule - collect three cores, escape north and press E - then reduces to one current instruction, a target direction and distance.")
    for text in [
        "WASD moves, mouse controls the camera and Space jumps.",
        "Shift sprints; C crouches; Q throws an echo stone; E interacts.",
        "A core needs a 1.5-second hold. Releasing E, leaving range or entering Chase interrupts it.",
        "At 3/3, the gate unlocks and points north; the player must press E at the gate to complete the run.",
    ]:
        bullet(doc, text)

    heading(doc, "5. Working prototype A: interaction, progress and persistence")
    paragraph(doc, "Prototype A combines PlayerInteractor, Collectible, GameManager, ExitGate, checkpoints and JSON saving. The interactor actively scans nearby colliders so interaction does not depend on a missed trigger callback. InteractionViewData publishes Available, Holding, Interrupted and Completed states to the HUD. Unique core IDs prevent duplicate scoring. Save version 3 retains checkpoint, core and relic progress, settings, run metrics and recent scores, with a safe fallback for damaged JSON.")

    heading(doc, "6. Working prototype B: character animation and stealth AI")
    paragraph(doc, "Prototype B combines camera-relative locomotion, semantic animation roles, baked NavMesh navigation and a guardian state machine. The player accelerates, brakes and rotates toward movement rather than sliding sideways. The PlayableGraph maps Idle, Walk, Run, Sprint, Crouch, Jump, Throw and Hit to a skeleton-bound Animator. Guardians use Patrol, Investigate, Search, Chase and Capture, with obstruction, suspicion, noise, shadow and crouch modifiers. Three separate patrol routes create observation, shadow and distraction challenges.")
    image(doc, ROOT / "Evidence" / "windows-gameplay.png", 4.6)
    p = paragraph(doc, "Figure 2. Production gameplay: player, cyan core target, warm environmental guidance and guardian threat are visible together.", after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    for run in p.runs:
        set_font(run, 8.5, MUTED, italic=True)

    heading(doc, "7. Testing and iteration")
    paragraph(doc, "Earlier videos revealed unreadable darkness, static character sliding, unclear E input and an invisible exit. These findings led directly to clarity settings, the Animator binding repair, hold-progress UI and the north-gate objective. Automated evidence is summarised below; three Windows playthroughs (no-alert, captured-then-complete and save-restore) remain the final manual recording check.")
    add_evidence_table(doc)

    heading(doc, "8. GitHub, assets and reflection")
    paragraph(doc, "Git records functional milestones from prototype systems through interaction, animation and verification. The public repository will contain code, used assets, settings, licence records and test evidence while excluding Library, Temp, Logs, Builds and recordings. External stealth projects informed state-machine organisation and feedback patterns only; no code, art or level content was copied. One refined map was prioritised over combat, networking and a second level to maximise reliability and evidence for the assessment.")

    heading(doc, "Appendix hand-in checklist")
    for text in [
        "Final screenshots: entrance, courtyard core, shadow gallery, echo-stone investigation, altar core, north-gate unlock and result screen.",
        "Automated results: TestResults-EditMode.xml and TestResults-PlayMode.xml.",
        "Three dated manual playthrough rows in PLAYTEST.md and one complete recording link.",
        "ASSET_LICENSES.md and the GitHub commit history.",
    ]:
        bullet(doc, text)

    doc.core_properties.title = "Echoes of the Ruins - Coursework Report"
    doc.core_properties.subject = "Unity 3D stealth exploration coursework"
    doc.core_properties.author = ""
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
